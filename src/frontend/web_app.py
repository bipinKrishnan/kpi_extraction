import os
import re
import time
import json
import hmac
import hashlib
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor

import pandas as pd
import requests
import streamlit as st
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


# ============================================================
# Streamlit / CommonMark markdown fix
# ============================================================

_st_markdown = st.markdown


def _html_safe_markdown(body, *args, **kwargs):
    if isinstance(body, str) and kwargs.get("unsafe_allow_html"):
        body = "\n".join(
            line.lstrip()
            for line in body.strip("\n").split("\n")
        )

    return _st_markdown(body, *args, **kwargs)


st.markdown = _html_safe_markdown


# ============================================================
# Configuration
# ============================================================

load_dotenv()

WORKSPACE_URL = os.getenv("WORKSPACE_URL")
MODAL_KEY = os.getenv("MODAL_KEY")
MODAL_SECRET = os.getenv("MODAL_SECRET")
APP_KEY = os.getenv("APP_KEY")

if not WORKSPACE_URL:
    st.error("WORKSPACE_URL is missing from environment variables.")
    st.stop()

if not MODAL_KEY:
    st.error("MODAL_KEY is missing from environment variables.")
    st.stop()

if not MODAL_SECRET:
    st.error("MODAL_SECRET is missing from environment variables.")
    st.stop()

if not APP_KEY:
    st.error(
        "APP_KEY is missing. "
        "Make sure the Modal secret 'kpi-streamlit-secrets' "
        "is attached to the Streamlit deployment."
    )
    st.stop()

ENDPOINT = f"{WORKSPACE_URL.rstrip('/')}/extract/text"

ALLOWED_EXTENSIONS = [
    "pdf",
    "pptx",
    "docx",
]

MAX_FILES = 20
MAX_FILE_SIZE_MB = 50
MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024
MAX_WORKERS = 5
REQUEST_START_DELAY_SECONDS = 2

MEASURE_ORDER = [
    "actual",
    "actual_ly",
    "budget",
    "forecast",
]

PERIOD_TYPE_ORDER = {
    "MTH": 1,
    "YTD": 2,
    "LTM": 3,
    "FY": 4,
}

FX_BASIS_OPTIONS = [
    "constant",
    "reported",
]

# ------------------------------------------------------------
# Metric display order for the Financial KPIs matrix.
#
# Metrics are ordered by this priority list first (based on
# availability — metrics not in this list are pushed to the
# end), then by increasing length of their "dimensions" list.
# ------------------------------------------------------------
METRIC_ORDER = [
    "revenue",
    "sales",
    "adjusted gross margin",
    "adjusted gross margin (w/o d&a)",
    "opex",
    "adjusted ebitda",
    "one-off",
    "reported ebitda",
    "capex",
    "cash conversion",
    "nfp",
]


def metric_sort_rank(metric_name):
    """
    Return the priority rank of a metric name for ordering
    purposes. Unrecognized metric names are pushed after all
    recognized ones (in the order METRIC_ORDER lists them).
    """

    if metric_name is None:
        return len(METRIC_ORDER)

    normalized = str(metric_name).strip().lower()

    if normalized in METRIC_ORDER:
        return METRIC_ORDER.index(normalized)

    return len(METRIC_ORDER)


XLSX_DIRECTORY = os.getenv(
    "XLSX_DIRECTORY",
    "/root/xlsx",
)

os.makedirs(XLSX_DIRECTORY, exist_ok=True)


# ============================================================
# Streamlit page configuration
# ============================================================

st.set_page_config(
    page_title="VerePort",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="collapsed",
)


# ============================================================
# Apple-inspired UI
# ============================================================

def inject_apple_theme():
    """
    Apple-inspired visual system for the entire Streamlit app.
    """

    st.session_state["ui_theme"] = "dark"
    theme = "dark"

    dark_css = """
    <style>

    /* =========================================================
       VerePort — Apple-inspired dark system
       ========================================================= */

    :root {
        color-scheme: dark;

        --vp-page: #000000;
        --vp-surface-1: #1d1d1f;
        --vp-surface-2: #232325;
        --vp-surface-3: #2d2d30;
        --vp-surface-hover: #212123;

        --vp-chrome: rgba(29,29,31,0.72);
        --vp-chrome-solid: rgba(29,29,31,0.92);

        --vp-text-primary: #f5f5f7;
        --vp-text-secondary: #a1a1a6;
        --vp-text-muted: #86868b;

        --vp-grid: #2a2a2c;
        --vp-axis: #424245;

        --vp-border: rgba(255,255,255,0.09);
        --vp-border-strong: rgba(255,255,255,0.18);

        --vp-fill-0: rgba(255,255,255,0.02);
        --vp-fill-1: rgba(255,255,255,0.045);
        --vp-fill-2: rgba(255,255,255,0.08);
        --vp-fill-3: rgba(255,255,255,0.16);

        --vp-accent: #2997ff;
        --vp-accent-solid: #0071e3;
        --vp-accent-hover: #0077ed;
        --vp-accent-tint: rgba(41,151,255,0.08);

        --vp-good: #30d158;
        --vp-warning: #ffd60a;
        --vp-serious: #ff9f0a;
        --vp-critical: #ff453a;

        --vp-sans:
            -apple-system,
            BlinkMacSystemFont,
            "SF Pro Display",
            "SF Pro Text",
            "Helvetica Neue",
            Helvetica,
            Arial,
            sans-serif;

        --vp-mono:
            "SF Mono",
            ui-monospace,
            SFMono-Regular,
            Menlo,
            Monaco,
            Consolas,
            monospace;

        --vp-radius: 18px;
        --vp-radius-sm: 12px;
        --vp-pill: 980px;

        --vp-shadow: none;

        --vp-shadow-soft:
            0 1px 2px rgba(0,0,0,0.25),
            0 10px 35px rgba(0,0,0,0.18);
    }


    /* =========================================================
       Base application
       ========================================================= */

    html,
    body,
    [data-testid="stAppViewContainer"],
    [data-testid="stApp"] {
        background: var(--vp-page) !important;
        color: var(--vp-text-primary) !important;
        font-family: var(--vp-sans) !important;
    }

    body {
        -webkit-font-smoothing: antialiased;
        text-rendering: optimizeLegibility;
    }

    .stApp {
        background:
            radial-gradient(
                circle at 50% -15%,
                rgba(41,151,255,0.07),
                transparent 34%
            ),
            var(--vp-page) !important;
    }

    [data-testid="stAppViewContainer"] {
        background: transparent !important;
        margin-left: 0 !important;
    }

    [data-testid="stHeader"] {
        background: transparent !important;
        height: 0 !important;
    }

    [data-testid="stToolbar"] {
        visibility: hidden;
        height: 0;
    }

    #MainMenu {
        visibility: hidden;
    }

    footer {
        visibility: hidden;
    }

    .block-container {
        max-width: 1480px;
        padding-top: 2.4rem !important;
        padding-bottom: 5rem !important;
        padding-left: 3rem !important;
        padding-right: 3rem !important;
    }


    /* =========================================================
       Typography
       ========================================================= */

    h1,
    h2,
    h3,
    h4,
    h5,
    h6 {
        font-family: var(--vp-sans) !important;
        color: var(--vp-text-primary) !important;
        letter-spacing: -0.025em !important;
    }

    h1 {
        font-size: clamp(2.8rem, 5vw, 4.8rem) !important;
        line-height: 0.98 !important;
        font-weight: 700 !important;
        letter-spacing: -0.055em !important;
        margin-bottom: 1rem !important;
    }

    h2 {
        font-size: 2rem !important;
        line-height: 1.08 !important;
        font-weight: 650 !important;
    }

    h3 {
        font-size: 1.45rem !important;
        line-height: 1.15 !important;
        font-weight: 650 !important;
    }

    p,
    li,
    label,
    .stMarkdown,
    .stCaption,
    .stText,
    .stAlert {
        font-family: var(--vp-sans) !important;
    }

    .stCaption,
    [data-testid="stCaptionContainer"] {
        color: var(--vp-text-secondary) !important;
        font-size: 0.92rem !important;
    }

    code,
    pre {
        font-family: var(--vp-mono) !important;
    }


    /* =========================================================
       Sidebar — completely removed
       ========================================================= */

    section[data-testid="stSidebar"],
    [data-testid="stSidebarCollapsedControl"],
    [data-testid="stSidebarNav"] {
        display: none !important;
        width: 0 !important;
        min-width: 0 !important;
        max-width: 0 !important;
    }


    /* =========================================================
       Top-right logout
       ========================================================= */

    [data-testid="stHorizontalBlock"] {
        align-items: flex-start !important;
    }

    div[data-testid="column"]:has(button[key="top_right_logout"]) {
        margin-left: auto !important;
    }

    button[key="top_right_logout"] {
        white-space: nowrap !important;
    }


    /* =========================================================
       Cards
       ========================================================= */

    .vp-card {
        padding: 1.55rem;

        border:
            1px solid var(--vp-border);

        border-radius: var(--vp-radius);

        background:
            linear-gradient(
                145deg,
                rgba(35,35,37,0.96),
                rgba(29,29,31,0.96)
            );

        box-shadow: var(--vp-shadow);

        margin-bottom: 1rem;
    }

    .vp-card-subtle {
        background: var(--vp-fill-1);
    }

    .vp-card-title {
        color: var(--vp-text-primary);
        font-size: 1.05rem;
        font-weight: 650;
        letter-spacing: -0.015em;
        margin-bottom: 0.35rem;
    }

    .vp-card-description {
        color: var(--vp-text-secondary);
        font-size: 0.9rem;
        line-height: 1.5;
    }


    /* =========================================================
       KPI metric cards
       ========================================================= */

    [data-testid="stMetric"] {
        padding: 1.35rem 1.45rem !important;

        border:
            1px solid var(--vp-border) !important;

        border-radius:
            var(--vp-radius) !important;

        background:
            linear-gradient(
                145deg,
                rgba(35,35,37,0.96),
                rgba(29,29,31,0.96)
            ) !important;

        box-shadow: none !important;
    }

    [data-testid="stMetricLabel"] {
        color: var(--vp-text-secondary) !important;
        font-size: 0.82rem !important;
        font-weight: 550 !important;
    }

    [data-testid="stMetricValue"] {
        color: var(--vp-text-primary) !important;
        font-size: 2rem !important;
        font-weight: 650 !important;
        letter-spacing: -0.04em !important;
    }


    /* =========================================================
       Buttons
       ========================================================= */

    .stButton > button,
    .stDownloadButton > button,
    button[kind="primary"] {
        min-height: 44px !important;

        padding:
            0.62rem 1.15rem !important;

        border-radius:
            var(--vp-pill) !important;

        border:
            1px solid rgba(255,255,255,0.12) !important;

        background:
            var(--vp-surface-3) !important;

        color:
            var(--vp-text-primary) !important;

        font-family:
            var(--vp-sans) !important;

        font-weight:
            600 !important;

        letter-spacing:
            -0.01em !important;

        transition:
            background 160ms ease,
            transform 160ms ease,
            border-color 160ms ease !important;

        box-shadow: none !important;
        cursor: pointer !important;
    }

    .stButton > button:hover,
    .stDownloadButton > button:hover {
        background:
            #363639 !important;

        border-color:
            var(--vp-border-strong) !important;

        transform:
            translateY(-1px);
    }

    .stButton > button:active,
    .stDownloadButton > button:active {
        background:
            #2e2e31 !important;

        transform:
            translateY(1px) scale(0.98);

        transition-duration:
            60ms;
    }

    button[kind="primary"],
    .stButton > button[kind="primary"],
    .stDownloadButton > button[kind="primary"] {
        background:
            linear-gradient(
                180deg,
                #2997ff,
                #0071e3
            ) !important;

        border-color:
            rgba(255,255,255,0.08) !important;

        color:
            white !important;

        box-shadow:
            0 6px 18px rgba(0,113,227,0.24) !important;
    }

    button[kind="primary"]:hover,
    .stButton > button[kind="primary"]:hover,
    .stDownloadButton > button[kind="primary"]:hover {
        background:
            linear-gradient(
                180deg,
                #40a4ff,
                #0077ed
            ) !important;

        box-shadow:
            0 8px 22px rgba(0,113,227,0.30) !important;
    }

    button[kind="primary"]:active,
    .stButton > button[kind="primary"]:active,
    .stDownloadButton > button[kind="primary"]:active {
        background:
            linear-gradient(
                180deg,
                #2490f4,
                #0067d1
            ) !important;

        transform:
            translateY(1px) scale(0.98);

        box-shadow:
            0 3px 10px rgba(0,113,227,0.22) !important;

        transition-duration:
            60ms;
    }


    /* =========================================================
       Inputs
       ========================================================= */

    .stTextInput input,
    .stSelectbox div[data-baseweb="select"] > div,
    .stMultiSelect div[data-baseweb="select"] > div,
    .stNumberInput input,
    .stDateInput input,
    .stTextArea textarea {
        min-height: 44px !important;

        border:
            1px solid var(--vp-border) !important;

        border-radius:
            var(--vp-radius-sm) !important;

        background:
            var(--vp-fill-1) !important;

        color:
            var(--vp-text-primary) !important;

        font-family:
            var(--vp-sans) !important;

        box-shadow:
            inset 0 1px 0 rgba(255,255,255,0.025) !important;
    }

    .stTextInput input:focus,
    .stNumberInput input:focus,
    .stTextArea textarea:focus {
        border-color:
            rgba(41,151,255,0.72) !important;

        box-shadow:
            0 0 0 3px rgba(41,151,255,0.12) !important;
    }

    .stTextInput label,
    .stSelectbox label,
    .stMultiSelect label,
    .stFileUploader label,
    .stDateInput label {
        color:
            var(--vp-text-secondary) !important;

        font-size:
            0.84rem !important;

        font-weight:
            550 !important;
    }


    /* =========================================================
       File uploader
       ========================================================= */

    [data-testid="stFileUploader"] {
        padding: 0.25rem;
    }

    [data-testid="stFileUploaderDropzone"] {
        min-height: 180px;

        border:
            1px dashed rgba(255,255,255,0.20) !important;

        border-radius:
            var(--vp-radius) !important;

        background:
            linear-gradient(
                145deg,
                rgba(255,255,255,0.045),
                rgba(255,255,255,0.025)
            ) !important;

        transition:
            border-color 160ms ease,
            background 160ms ease !important;
    }

    [data-testid="stFileUploaderDropzone"]:hover {
        border-color:
            rgba(41,151,255,0.65) !important;

        background:
            rgba(41,151,255,0.045) !important;
    }

    [data-testid="stFileUploaderDropzoneInstructions"] {
        color:
            var(--vp-text-secondary) !important;
    }

    [data-testid="stFileUploaderDropzoneInstructions"] small {
        color:
            var(--vp-text-muted) !important;
    }


    /* =========================================================
       Dataframes
       ========================================================= */

    [data-testid="stDataFrame"] {
        border:
            1px solid var(--vp-border) !important;

        border-radius:
            var(--vp-radius) !important;

        overflow:
            hidden !important;

        background:
            var(--vp-surface-1) !important;
    }


    /* =========================================================
       Alerts
       ========================================================= */

    [data-testid="stAlert"] {
        border-radius:
            var(--vp-radius-sm) !important;

        border:
            1px solid var(--vp-border) !important;

        background:
            var(--vp-fill-1) !important;

        color:
            var(--vp-text-secondary) !important;
    }


    /* =========================================================
       Expanders
       ========================================================= */

    [data-testid="stExpander"] {
        border:
            1px solid var(--vp-border) !important;

        border-radius:
            var(--vp-radius) !important;

        background:
            var(--vp-fill-1) !important;

        overflow:
            hidden !important;
    }

    [data-testid="stExpander"] summary {
        color:
            var(--vp-text-primary) !important;

        font-weight:
            600 !important;
    }


    /* =========================================================
       Progress
       ========================================================= */

    [data-testid="stProgressBar"] {
        padding-top: 0.4rem;
    }

    [data-testid="stProgressBar"] > div {
        background:
            rgba(255,255,255,0.10) !important;

        border-radius:
            var(--vp-pill) !important;
    }

    [data-testid="stProgressBar"] > div > div {
        background:
            linear-gradient(
                90deg,
                #2997ff,
                #0071e3
            ) !important;

        border-radius:
            var(--vp-pill) !important;
    }


    /* =========================================================
       Dividers
       ========================================================= */

    hr {
        border:
            none !important;

        border-top:
            1px solid var(--vp-border) !important;

        margin:
            2rem 0 !important;
    }


    /* =========================================================
       Login
       ========================================================= */

    /*
       The title and form are now inside the same Streamlit
       container. This keeps the existing login content exactly
       as-is while allowing the card to position both elements
       together.
    */

    .st-key-login_container {
        width: min(520px, 100%) !important;
        margin: 12vh auto 0 auto !important;

        padding: 2.5rem !important;

        border:
            1px solid var(--vp-border) !important;

        border-radius:
            28px !important;

        background:
            linear-gradient(
                145deg,
                rgba(35,35,37,0.98),
                rgba(18,18,19,0.98)
            ) !important;

        box-shadow:
            0 24px 70px rgba(0,0,0,0.38) !important;
    }

    .vp-login-title {
        color: var(--vp-text-primary);
        font-size: 2.6rem;
        line-height: 1;
        font-weight: 700;
        letter-spacing: -0.05em;
        margin: 0 0 1.5rem 0;
    }

    .st-key-login_container [data-testid="stForm"] {
        border: none !important;
        padding: 0 !important;
        background: transparent !important;
    }

    .st-key-login_container [data-testid="stFormSubmitButton"] button {
        margin-top: 0.35rem !important;
    }


    /* =========================================================
       Extraction result section
       ========================================================= */

    .vp-result-banner {
        padding: 1.2rem 1.35rem;

        border:
            1px solid rgba(48,209,88,0.20);

        border-radius:
            var(--vp-radius);

        background:
            rgba(48,209,88,0.055);

        margin-bottom:
            1.2rem;
    }

    .vp-result-title {
        color: var(--vp-text-primary);
        font-weight: 650;
    }

    .vp-result-copy {
        color: var(--vp-text-secondary);
        font-size: 0.88rem;
        margin-top: 0.25rem;
    }


    /* =========================================================
       Small metadata rows
       ========================================================= */

    .vp-meta-row {
        display: flex;
        flex-wrap: wrap;
        gap: 8px;

        margin-top: 0.8rem;
    }

    .vp-meta-item,
    .vp-file-chip {
        padding: 6px 10px;

        border:
            1px solid var(--vp-border);

        border-radius:
            var(--vp-pill);

        background:
            var(--vp-fill-1);

        color:
            var(--vp-text-secondary);

        font-size:
            0.77rem;
    }


    /* =========================================================
       Footer
       ========================================================= */

    .vp-footer {
        padding-top: 2rem;

        color:
            var(--vp-text-muted);

        font-size:
            0.75rem;

        text-align:
            center;
    }


    /* =========================================================
       Mobile
       ========================================================= */

    @media (max-width: 900px) {

        .block-container {
            padding-left: 1rem !important;
            padding-right: 1rem !important;
            padding-top: 1.25rem !important;
        }

        .st-key-login_container {
            margin-top: 7vh !important;
            padding: 2rem 1.4rem !important;
            border-radius: 22px !important;
        }
    }

    </style>
    """

    light_css = """
    <style>

    :root {
        color-scheme: light;

        --vp-page: #f5f5f7;
        --vp-surface-1: #ffffff;
        --vp-surface-2: #f5f5f7;
        --vp-surface-3: #ebebed;
        --vp-surface-hover: #fafafc;

        --vp-text-primary: #1d1d1f;
        --vp-text-secondary: #6e6e73;
        --vp-text-muted: #86868b;

        --vp-border: rgba(0,0,0,0.08);
        --vp-border-strong: rgba(0,0,0,0.16);

        --vp-fill-1: rgba(0,0,0,0.035);

        --vp-sans:
            -apple-system,
            BlinkMacSystemFont,
            "SF Pro Display",
            "SF Pro Text",
            "Helvetica Neue",
            Helvetica,
            Arial,
            sans-serif;

        --vp-radius: 18px;
        --vp-radius-sm: 12px;
        --vp-pill: 980px;
    }

    body,
    [data-testid="stAppViewContainer"],
    [data-testid="stApp"] {
        background: var(--vp-page) !important;
        color: var(--vp-text-primary) !important;
    }

    .stApp {
        background:
            radial-gradient(
                circle at 50% -15%,
                rgba(0,113,227,0.055),
                transparent 34%
            ),
            var(--vp-page) !important;
    }

    [data-testid="stAppViewContainer"] {
        background: transparent !important;
    }

    [data-testid="stHeader"] {
        background: transparent !important;
        height: 0 !important;
    }

    [data-testid="stToolbar"] {
        visibility: hidden;
        height: 0;
    }

    #MainMenu {
        visibility: hidden;
    }

    footer {
        visibility: hidden;
    }

    .block-container {
        max-width: 1480px;
        padding-top: 2.4rem !important;
        padding-bottom: 5rem !important;
        padding-left: 3rem !important;
        padding-right: 3rem !important;
    }

    h1,
    h2,
    h3,
    h4,
    h5,
    h6 {
        color: var(--vp-text-primary) !important;
    }

    h1 {
        font-size: clamp(2.8rem, 5vw, 4.8rem) !important;
        line-height: 0.98 !important;
        font-weight: 700 !important;
        letter-spacing: -0.055em !important;
    }

    h2 {
        font-size: 2rem !important;
        line-height: 1.08 !important;
        font-weight: 650 !important;
    }

    h3 {
        font-size: 1.45rem !important;
        line-height: 1.15 !important;
        font-weight: 650 !important;
    }

    p,
    li,
    label,
    .stMarkdown,
    .stCaption,
    .stText {
        font-family: var(--vp-sans) !important;
    }

    .stCaption,
    [data-testid="stCaptionContainer"] {
        color: var(--vp-text-secondary) !important;
    }

    section[data-testid="stSidebar"],
    [data-testid="stSidebarCollapsedControl"],
    [data-testid="stSidebarNav"] {
        display: none !important;
        width: 0 !important;
        min-width: 0 !important;
        max-width: 0 !important;
    }

    .vp-card {
        background:
            rgba(255,255,255,0.92);

        border-color:
            var(--vp-border);

        box-shadow:
            0 1px 2px rgba(0,0,0,0.04),
            0 8px 24px rgba(0,0,0,0.06);
    }

    [data-testid="stMetric"] {
        background:
            rgba(255,255,255,0.92) !important;

        border-color:
            var(--vp-border) !important;

        box-shadow:
            0 1px 2px rgba(0,0,0,0.04),
            0 8px 24px rgba(0,0,0,0.06) !important;
    }

    [data-testid="stMetricLabel"] {
        color:
            var(--vp-text-secondary) !important;
    }

    [data-testid="stMetricValue"] {
        color:
            var(--vp-text-primary) !important;
    }

    .stButton > button,
    .stDownloadButton > button,
    button[kind="primary"] {
        border-color:
            var(--vp-border) !important;

        background:
            #ffffff !important;

        color:
            var(--vp-text-primary) !important;

        box-shadow:
            0 1px 2px rgba(0,0,0,0.04) !important;
    }

    .stButton > button:hover,
    .stDownloadButton > button:hover {
        background:
            #f5f5f7 !important;

        border-color:
            var(--vp-border-strong) !important;
    }

    .stButton > button:active,
    .stDownloadButton > button:active {
        background:
            #e8e8ed !important;

        transform:
            translateY(1px) scale(0.98);

        transition-duration:
            60ms;
    }

    button[kind="primary"],
    .stButton > button[kind="primary"],
    .stDownloadButton > button[kind="primary"] {
        background:
            linear-gradient(
                180deg,
                #147ce5,
                #0071e3
            ) !important;

        color:
            #ffffff !important;

        border-color:
            rgba(0,0,0,0.03) !important;

        box-shadow:
            0 6px 18px rgba(0,113,227,0.18) !important;
    }

    button[kind="primary"]:active,
    .stButton > button[kind="primary"]:active,
    .stDownloadButton > button[kind="primary"]:active {
        background:
            linear-gradient(
                180deg,
                #1272cf,
                #0062c4
            ) !important;

        transform:
            translateY(1px) scale(0.98);

        box-shadow:
            0 2px 8px rgba(0,113,227,0.18) !important;

        transition-duration:
            60ms;
    }

    .stTextInput input,
    .stSelectbox div[data-baseweb="select"] > div,
    .stMultiSelect div[data-baseweb="select"] > div,
    .stNumberInput input,
    .stDateInput input,
    .stTextArea textarea {
        background:
            rgba(255,255,255,0.92) !important;

        color:
            var(--vp-text-primary) !important;

        border-color:
            var(--vp-border) !important;
    }

    [data-testid="stFileUploaderDropzone"] {
        border-color:
            rgba(0,0,0,0.16) !important;

        background:
            rgba(255,255,255,0.68) !important;
    }

    [data-testid="stFileUploaderDropzone"]:hover {
        border-color:
            rgba(0,113,227,0.55) !important;

        background:
            rgba(0,113,227,0.035) !important;
    }

    [data-testid="stDataFrame"] {
        background:
            #ffffff !important;

        border-color:
            var(--vp-border) !important;
    }

    [data-testid="stAlert"] {
        background:
            rgba(255,255,255,0.80) !important;

        border-color:
            var(--vp-border) !important;
    }

    [data-testid="stExpander"] {
        background:
            rgba(255,255,255,0.78) !important;

        border-color:
            var(--vp-border) !important;
    }

    [data-testid="stProgressBar"] > div {
        background:
            rgba(0,0,0,0.08) !important;
    }

    [data-testid="stProgressBar"] > div > div {
        background:
            linear-gradient(
                90deg,
                #2a8cff,
                #0071e3
            ) !important;
    }

    hr {
        border-top:
            1px solid var(--vp-border) !important;
    }

    .vp-result-banner {
        background:
            rgba(26,127,55,0.055);

        border-color:
            rgba(26,127,55,0.16);
    }

    /* =========================================================
       Login
       ========================================================= */

    .st-key-login_container {
        width: min(520px, 100%) !important;
        margin: 12vh auto 0 auto !important;

        padding: 2.5rem !important;

        border:
            1px solid var(--vp-border) !important;

        border-radius:
            28px !important;

        background:
            rgba(255,255,255,0.96) !important;

        box-shadow:
            0 24px 70px rgba(0,0,0,0.10) !important;
    }

    .vp-login-title {
        color: var(--vp-text-primary);
        font-size: 2.6rem;
        line-height: 1;
        font-weight: 700;
        letter-spacing: -0.05em;
        margin: 0 0 1.5rem 0;
    }

    .st-key-login_container [data-testid="stForm"] {
        border: none !important;
        padding: 0 !important;
        background: transparent !important;
    }

    .st-key-login_container [data-testid="stFormSubmitButton"] button {
        margin-top: 0.35rem !important;
    }

    @media (max-width: 900px) {

        .block-container {
            padding-left: 1rem !important;
            padding-right: 1rem !important;
            padding-top: 1.25rem !important;
        }

        .st-key-login_container {
            margin-top: 7vh !important;
            padding: 2rem 1.4rem !important;
            border-radius: 22px !important;
        }
    }

    </style>
    """

    st.markdown(
        dark_css if theme == "dark" else light_css,
        unsafe_allow_html=True,
    )


inject_apple_theme()


# ============================================================
# Authentication
# ============================================================

def authenticate_app():
    """
    Display the login screen and authenticate the user against
    APP_KEY supplied through the Modal secret.
    """

    if "authenticated" not in st.session_state:
        st.session_state["authenticated"] = False

    if st.session_state["authenticated"]:

        logout_col = st.columns(
            [1, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1]
        )[11]

        with logout_col:
            if st.button(
                "Log out",
                key="top_right_logout",
                width="stretch",
            ):
                st.session_state["authenticated"] = False
                st.session_state.pop("result_df", None)
                st.session_state.pop(
                    "conflict_resolutions",
                    None,
                )
                st.rerun()

        return True

    # The title and form are deliberately inside the same
    # Streamlit container so the card is positioned correctly.
    with st.container(key="login_container"):

        st.markdown(
            """
            <div class="vp-login-title">
                VerePort Extract
            </div>
            """,
            unsafe_allow_html=True,
        )

        with st.form(
            "login_form",
            clear_on_submit=False,
        ):

            entered_app_key = st.text_input(
                "Application Key",
                type="password",
                placeholder="Enter your application key",
                help="Enter the application key provided to you.",
            )

            submitted = st.form_submit_button(
                "Continue",
                type="primary",
                width="stretch",
            )

        if submitted:

            entered_app_key = entered_app_key.strip()

            if not entered_app_key:
                st.error("Please enter an application key.")
                return False

            if hmac.compare_digest(
                entered_app_key,
                APP_KEY,
            ):
                st.session_state["authenticated"] = True
                st.rerun()

            else:
                st.error("The application key is not valid.")

    return False


# ============================================================
# Authentication Gate
# ============================================================

if not authenticate_app():
    st.stop()


# ============================================================
# Generic helpers
# ============================================================

def is_missing(value):
    """Safely determine whether a value should be treated as missing."""

    if value is None:
        return True

    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def clean_value(value):
    """Return None for missing values, otherwise preserve the value."""

    if is_missing(value):
        return None

    if isinstance(value, str):
        value = value.strip()

        if not value:
            return None

    return value


# ============================================================
# Streamlit display sanitization
# ============================================================

def sanitize_for_streamlit(df):
    """
    Ensure every column has a single, Arrow-compatible dtype
    before rendering with st.dataframe.

    Streamlit converts DataFrames to Arrow tables internally.
    Arrow requires a consistent type per column, but our matrix
    and detail tables can contain a mix of numbers, strings, and
    None within the same column (e.g. unresolved conflict text
    alongside numeric values). Without this, st.dataframe raises
    pyarrow.lib.ArrowTypeError.
    """

    sanitized = df.copy()

    for column in sanitized.columns:

        if sanitized[column].dtype == object:

            sanitized[column] = sanitized[column].apply(
                lambda value: (
                    ""
                    if is_missing(value)
                    else str(value)
                )
            )

    return sanitized


def normalize_scalar(value):
    """
    Convert API fields that may arrive as one-item lists/tuples
    into scalars.
    """

    if isinstance(value, (list, tuple)):

        if len(value) == 0:
            return None

        if len(value) == 1:
            return value[0]

        return ", ".join(
            str(item)
            for item in value
        )

    return value


# ============================================================
# Dimension helpers
# ============================================================

def normalize_dimensions(dimensions):
    """Normalize the API dimensions field into a list of strings."""

    if dimensions is None:
        return []

    if isinstance(dimensions, list):
        return [
            str(item)
            for item in dimensions
            if item is not None
        ]

    if isinstance(dimensions, tuple):
        return [
            str(item)
            for item in dimensions
            if item is not None
        ]

    if isinstance(dimensions, dict):
        return [
            str(value)
            for value in dimensions.values()
            if value is not None
        ]

    return [str(dimensions)]


def format_dimensions(dimensions):
    """Keep dimensions represented as JSON-style list in Excel."""

    normalized = normalize_dimensions(dimensions)

    return json.dumps(
        normalized,
        ensure_ascii=False,
    )


# ============================================================
# Page helper
# ============================================================

def extract_page_number(label):
    """Extract page information from the API extraction label."""

    if label is None:
        return None

    label_text = str(label).strip()

    match = re.search(
        r"\bpages?\s*[:=_\-]?\s*(\d+(?:\s*[-–]\s*\d+)?)",
        label_text,
        flags=re.IGNORECASE,
    )

    if match:
        return match.group(1)

    match = re.search(
        r"\bp\.?\s*(\d+(?:\s*[-–]\s*\d+)?)\b",
        label_text,
        flags=re.IGNORECASE,
    )

    if match:
        return match.group(1)

    return label_text


# ============================================================
# Value helper
# ============================================================

def convert_value(value):
    """Convert numeric-looking API values to Python numbers."""

    value = clean_value(value)

    if value is None:
        return None

    if isinstance(value, bool):
        return value

    if isinstance(value, (int, float)):
        return value

    text = str(value).strip()

    numeric_text = text.replace(",", "").strip()

    if numeric_text.endswith("%"):

        try:
            number = float(
                numeric_text[:-1].strip()
            )

            return number / 100

        except ValueError:
            return value

    try:

        if re.fullmatch(
            r"[-+]?\d+",
            numeric_text,
        ):
            return int(numeric_text)

        if re.fullmatch(
            r"[-+]?(?:\d+\.\d*|\.\d+)",
            numeric_text,
        ):
            return float(numeric_text)

    except ValueError:
        pass

    return value


# ============================================================
# Period helpers
# ============================================================

def normalize_period_type(period_type):
    """Normalize period type to uppercase."""

    period_type = normalize_scalar(period_type)

    if period_type is None:
        return None

    return str(
        period_type
    ).strip().upper()


def normalize_measure(measure):
    """Normalize measure to lowercase."""

    measure = normalize_scalar(measure)

    if measure is None:
        return None

    return str(
        measure
    ).strip().lower()


def normalize_fx_basis(fx_basis):
    """Normalize FX basis into a readable scalar."""

    fx_basis = normalize_scalar(fx_basis)

    if fx_basis is None:
        return None

    return str(
        fx_basis
    ).strip()


def period_sort_key(period):
    """Create a useful chronological sort key."""

    if period is None:
        return (
            9999,
            99,
            99,
            "",
        )

    text = str(
        period
    ).strip().upper()

    match = re.search(
        r"\b(20\d{2})[-/](\d{1,2})\b",
        text,
    )

    if match:
        return (
            int(match.group(1)),
            int(match.group(2)),
            0,
            text,
        )

    match = re.search(
        r"\bQ([1-4])\s*[-/]?\s*(20\d{2})\b",
        text,
    )

    if match:
        return (
            int(match.group(2)),
            int(match.group(1)) * 3,
            1,
            text,
        )

    match = re.search(
        r"\b(20\d{2})\s*[-/]?\s*Q([1-4])\b",
        text,
    )

    if match:
        return (
            int(match.group(1)),
            int(match.group(2)) * 3,
            1,
            text,
        )

    match = re.search(
        r"\bFY\s*(20\d{2})\b",
        text,
    )

    if match:
        return (
            int(match.group(1)),
            12,
            2,
            text,
        )

    match = re.fullmatch(
        r"(20\d{2})",
        text,
    )

    if match:
        return (
            int(match.group(1)),
            12,
            2,
            text,
        )

    return (
        9999,
        99,
        99,
        text,
    )


def column_sort_key(column):
    """Sort financial matrix columns chronologically."""

    period = column["period"]
    period_type = column["period_type"]
    measure = column["measure"]

    return (
        period_sort_key(period),
        PERIOD_TYPE_ORDER.get(
            period_type,
            99,
        ),
        (
            MEASURE_ORDER.index(measure)
            if measure in MEASURE_ORDER
            else 99
        ),
        str(
            measure or ""
        ),
    )


# ============================================================
# Flatten API response
# ============================================================

def get_extractions(response_data):
    """Support both extraction wrapper and direct KPI response."""

    if not isinstance(
        response_data,
        dict,
    ):
        return []

    if isinstance(
        response_data.get("extractions"),
        list,
    ):
        return (
            response_data.get("extractions")
            or []
        )

    if isinstance(
        response_data.get("kpis"),
        list,
    ):
        return [
            {
                "source": None,
                "label": None,
                "data": {
                    "kpis": (
                        response_data.get("kpis")
                        or []
                    ),
                },
            }
        ]

    return []


def flatten_extractions(
    response_data,
    default_source,
):
    """Convert backend response into one row per KPI observation."""

    rows = []

    if not isinstance(
        response_data,
        dict,
    ):
        return pd.DataFrame()

    extractions = get_extractions(
        response_data
    )

    for extraction in extractions:

        if not isinstance(
            extraction,
            dict,
        ):
            continue

        extraction_source = (
            extraction.get("source")
            or default_source
        )

        extraction_label = extraction.get(
            "label"
        )

        page_number = extract_page_number(
            extraction_label
        )

        data = (
            extraction.get("data")
            or {}
        )

        if not isinstance(
            data,
            dict,
        ):
            continue

        document = (
            data.get("document")
            or {}
        )

        if not isinstance(
            document,
            dict,
        ):
            document = {}

        company = document.get(
            "company"
        )

        holding = document.get(
            "holding"
        )

        reporting_period = document.get(
            "reporting_period"
        )

        document_source = document.get(
            "source"
        )

        source = (
            document_source
            or extraction_source
            or default_source
        )

        source_document = source

        kpis = (
            data.get("kpis")
            or []
        )

        if not isinstance(
            kpis,
            list,
        ):
            continue

        for kpi in kpis:

            if not isinstance(
                kpi,
                dict,
            ):
                continue

            kpi_name = kpi.get(
                "name"
            )

            observations = (
                kpi.get("observations")
                or []
            )

            if not isinstance(
                observations,
                list,
            ):
                continue

            for observation in observations:

                if not isinstance(
                    observation,
                    dict,
                ):
                    continue

                dimensions = normalize_dimensions(
                    observation.get(
                        "dimensions"
                    )
                )

                period = normalize_scalar(
                    observation.get(
                        "period"
                    )
                )

                period_type = normalize_period_type(
                    observation.get(
                        "period_type"
                    )
                )

                measure = normalize_measure(
                    observation.get(
                        "measure"
                    )
                )

                value = convert_value(
                    observation.get(
                        "value"
                    )
                )

                unit = normalize_scalar(
                    observation.get(
                        "unit"
                    )
                )

                currency = normalize_scalar(
                    observation.get(
                        "currency"
                    )
                )

                fx_basis = normalize_fx_basis(
                    observation.get(
                        "fx_basis"
                    )
                )

                rows.append(
                    {
                        "source": source,
                        "source_document": source_document,
                        "page_number": page_number,
                        "extraction_source": extraction_source,
                        "extraction_label": extraction_label,
                        "company": company,
                        "holding": holding,
                        "reporting_period": reporting_period,
                        "kpi": clean_value(kpi_name),
                        "measure": measure,
                        "period": period,
                        "period_type": period_type,
                        "value": value,
                        "currency": currency,
                        "fx_basis": fx_basis,
                        "unit": unit,
                        "dimensions": dimensions,
                    }
                )

    if not rows:
        return pd.DataFrame()

    return pd.DataFrame(rows)


# ============================================================
# Extract a single file
# ============================================================

def extract_single_file(
    uploaded_file,
    headers,
):
    """Send exactly one file in exactly one HTTP request."""

    filename = uploaded_file.name

    try:

        file_bytes = uploaded_file.getvalue()

        files = {
            "files": (
                filename,
                file_bytes,
                uploaded_file.type
                or "application/octet-stream",
            )
        }

        response = requests.post(
            ENDPOINT,
            headers=headers,
            files=files,
            timeout=3600,
        )

        response.raise_for_status()

        return {
            "filename": filename,
            "success": True,
            "response_data": response.json(),
            "error": None,
        }

    except requests.exceptions.Timeout:

        return {
            "filename": filename,
            "success": False,
            "response_data": None,
            "error": "Request timed out after 1 hour.",
        }

    except requests.exceptions.HTTPError as e:

        status_code = (
            e.response.status_code
            if e.response is not None
            else "unknown"
        )

        response_text = (
            e.response.text
            if e.response is not None
            else ""
        )

        return {
            "filename": filename,
            "success": False,
            "response_data": None,
            "error": (
                f"HTTP {status_code}"
                + (
                    f": {response_text}"
                    if response_text
                    else ""
                )
            ),
        }

    except requests.exceptions.RequestException as e:

        return {
            "filename": filename,
            "success": False,
            "response_data": None,
            "error": f"Request failed: {e}",
        }

    except ValueError:

        return {
            "filename": filename,
            "success": False,
            "response_data": None,
            "error": "The API returned invalid JSON.",
        }

    except Exception as e:

        return {
            "filename": filename,
            "success": False,
            "response_data": None,
            "error": f"Unexpected error: {e}",
        }


# ============================================================
# Financial matrix preparation
# ============================================================

def build_financial_matrix(df):

    if df.empty:
        return pd.DataFrame(), []

    working = df.copy()

    for column in [
        "kpi",
        "dimensions",
        "fx_basis",
        "unit",
        "currency",
        "period",
        "period_type",
        "measure",
        "value",
    ]:

        if column not in working.columns:
            working[column] = None

    working["kpi"] = working["kpi"].apply(
        lambda x: (
            str(x).strip()
            if not is_missing(x)
            else "Unnamed KPI"
        )
    )

    working["dimensions"] = working[
        "dimensions"
    ].apply(
        normalize_dimensions
    )

    working["dimensions_display"] = working[
        "dimensions"
    ].apply(
        format_dimensions
    )

    # Length of the dimensions list, used for the "increasing
    # length of dimensions" ordering requirement.
    working["dimensions_len"] = working[
        "dimensions"
    ].apply(len)

    working["fx_basis"] = working[
        "fx_basis"
    ].apply(
        lambda x: (
            str(x).strip()
            if not is_missing(x)
            else ""
        )
    )

    working["unit"] = working[
        "unit"
    ].apply(
        lambda x: (
            str(x).strip()
            if not is_missing(x)
            else ""
        )
    )

    working["period"] = working[
        "period"
    ].apply(
        normalize_scalar
    )

    working["period_type"] = working[
        "period_type"
    ].apply(
        normalize_period_type
    )

    working["measure"] = working[
        "measure"
    ].apply(
        normalize_measure
    )

    working = working[
        working["period"].notna()
        & working["period_type"].notna()
        & working["measure"].notna()
    ].copy()

    if working.empty:
        return pd.DataFrame(), []

    column_records = (
        working[
            [
                "period",
                "period_type",
                "measure",
            ]
        ]
        .drop_duplicates()
        .to_dict("records")
    )

    column_records.sort(
        key=column_sort_key
    )

    working["row_key"] = list(
        zip(
            working["kpi"],
            working["dimensions_display"],
            working["fx_basis"],
            working["unit"],
        )
    )

    grouped_values = (
        working.groupby(
            [
                "row_key",
                "period",
                "period_type",
                "measure",
            ],
            sort=False,
            dropna=False,
        )["value"]
        .apply(
            lambda values: collapse_duplicate_pipe_values(
                values.iloc[0]
                if len(values) == 1
                else " | ".join(
                    str(v)
                    for v in values
                )
            )
        )
        .reset_index()
    )

    matrix_rows = []
    row_sort_keys = []

    unique_rows = (
        working[
            [
                "row_key",
                "kpi",
                "dimensions_display",
                "dimensions_len",
                "fx_basis",
                "unit",
            ]
        ]
        .drop_duplicates(
            subset=["row_key"]
        )
    )

    for _, base_row in unique_rows.iterrows():

        row = {
            "Metric": base_row["kpi"],
            "Dimensions": base_row["dimensions_display"],
            "FX Basis": base_row["fx_basis"],
            "UoM": base_row["unit"],
        }

        row_key = base_row["row_key"]

        matching_values = grouped_values[
            grouped_values["row_key"] == row_key
        ]

        for column in column_records:

            period = column["period"]
            period_type = column["period_type"]
            measure = column["measure"]

            match = matching_values[
                (
                    matching_values["period"] == period
                )
                & (
                    matching_values["period_type"]
                    == period_type
                )
                & (
                    matching_values["measure"]
                    == measure
                )
            ]

            if match.empty:
                value = None
            else:
                value = match.iloc[0]["value"]
                if (
                    isinstance(value, (int, float))
                    and not isinstance(value, bool)
                    and value == 0
                ):
                    value = None

            row[
                (
                    period,
                    period_type,
                    measure,
                )
            ] = value

        matrix_rows.append(row)

        metric_rank = metric_sort_rank(
            base_row["kpi"]
        )

        # For metrics not in the priority list, group them
        # alphabetically by name before applying the
        # dimensions-length ordering.
        unmatched_name_key = (
            str(base_row["kpi"]).strip().lower()
            if metric_rank == len(METRIC_ORDER)
            else ""
        )

        row_sort_keys.append(
            (
                metric_rank,
                unmatched_name_key,
                base_row["dimensions_len"],
                str(base_row["fx_basis"] or "").lower(),
                str(base_row["unit"] or "").lower(),
            )
        )

    if matrix_rows:

        ordered = sorted(
            zip(row_sort_keys, matrix_rows),
            key=lambda item: item[0],
        )

        matrix_rows = [
            item[1]
            for item in ordered
        ]

    matrix_df = pd.DataFrame(
        matrix_rows
    )

    return (
        matrix_df,
        column_records,
    )


# ============================================================
# Excel formatting helpers
# ============================================================

def format_measure_label(measure):

    labels = {
        "actual": "Actual",
        "actual_ly": "Actual LY",
        "budget": "Budget",
        "forecast": "Forecast",
    }

    return labels.get(
        measure,
        str(measure).replace(
            "_",
            " ",
        ).title(),
    )


def format_period_header(
    period,
    period_type,
):
    return f"{period} – {period_type}"


def apply_value_number_format(
    cell,
    unit=None,
    value=None,
):

    if value is None:
        return

    unit_text = (
        str(unit).strip().lower()
        if unit is not None
        else ""
    )

    if "%" in unit_text or "percent" in unit_text:

        cell.number_format = "0.0%"

    elif any(
        token in unit_text
        for token in [
            "eur",
            "usd",
            "gbp",
            "chf",
            "jpy",
        ]
    ):

        cell.number_format = "#,##0.00"

    else:

        cell.number_format = "#,##0.00"


# ============================================================
# Create Financial KPIs worksheet
# ============================================================

def create_financial_kpis_sheet(
    workbook,
    result_df,
):

    worksheet = workbook.create_sheet(
        "Financial KPIs"
    )

    matrix_df, column_records = (
        build_financial_matrix(result_df)
    )

    title_font = Font(
        bold=True,
        size=16,
    )

    metadata_label_font = Font(
        bold=True,
        size=10,
    )

    metadata_value_font = Font(
        size=10,
    )

    header_font = Font(
        bold=True,
        size=10,
    )

    metric_font = Font(
        bold=True,
        size=10,
    )

    thin_gray = Side(
        style="thin",
        color="D9D9D9",
    )

    medium_gray = Side(
        style="medium",
        color="A6A6A6",
    )

    border = Border(
        bottom=thin_gray,
    )

    header_border = Border(
        top=medium_gray,
        bottom=medium_gray,
    )

    period_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    measure_fill = PatternFill(
        fill_type="solid",
        fgColor="EAF2F8",
    )

    metadata_row = (
        result_df.iloc[0]
        if not result_df.empty
        else {}
    )

    company = (
        metadata_row.get("company")
        if hasattr(metadata_row, "get")
        else None
    )

    holding = (
        metadata_row.get("holding")
        if hasattr(metadata_row, "get")
        else None
    )

    reporting_period = (
        metadata_row.get("reporting_period")
        if hasattr(metadata_row, "get")
        else None
    )

    worksheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=max(
            4 + len(column_records),
            4,
        ),
    )

    title_cell = worksheet.cell(
        row=1,
        column=1,
        value="Financial KPIs",
    )

    title_cell.font = title_font

    title_cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    worksheet.row_dimensions[1].height = 24

    metadata_items = []

    if (
        not is_missing(company)
        and str(company).strip()
    ):
        metadata_items.append(
            ("Company", company)
        )

    if (
        not is_missing(holding)
        and str(holding).strip()
    ):
        metadata_items.append(
            ("Holding", holding)
        )

    if (
        not is_missing(reporting_period)
        and str(reporting_period).strip()
    ):
        metadata_items.append(
            (
                "Reporting Period",
                reporting_period,
            )
        )

    metadata_row_number = 3

    for label, value in metadata_items:

        worksheet.cell(
            row=metadata_row_number,
            column=1,
            value=label,
        ).font = metadata_label_font

        worksheet.cell(
            row=metadata_row_number,
            column=2,
            value=value,
        ).font = metadata_value_font

        metadata_row_number += 1

    header_start_row = metadata_row_number + 1

    static_headers = [
        "Metric",
        "Dimensions",
        "FX Basis",
        "UoM",
    ]

    for column_index, header_name in enumerate(
        static_headers,
        start=1,
    ):

        worksheet.cell(
            row=header_start_row,
            column=column_index,
            value=header_name,
        )

        worksheet.merge_cells(
            start_row=header_start_row,
            start_column=column_index,
            end_row=header_start_row + 1,
            end_column=column_index,
        )

        cell = worksheet.cell(
            row=header_start_row,
            column=column_index,
        )

        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        cell.fill = period_fill
        cell.border = header_border

    period_groups = []

    group_start = 5
    previous_group = None

    for index, column in enumerate(
        column_records
    ):

        group = (
            column["period"],
            column["period_type"],
        )

        excel_column = 5 + index

        if previous_group is None:

            previous_group = group
            group_start = excel_column

        elif group != previous_group:

            period_groups.append(
                (
                    previous_group,
                    group_start,
                    excel_column - 1,
                )
            )

            previous_group = group
            group_start = excel_column

    if previous_group is not None:

        period_groups.append(
            (
                previous_group,
                group_start,
                4 + len(column_records),
            )
        )

    for group, start_column, end_column in period_groups:

        period, period_type = group

        worksheet.merge_cells(
            start_row=header_start_row,
            start_column=start_column,
            end_row=header_start_row,
            end_column=end_column,
        )

        cell = worksheet.cell(
            row=header_start_row,
            column=start_column,
            value=format_period_header(
                period,
                period_type,
            ),
        )

        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        cell.fill = period_fill
        cell.border = header_border

        for column in range(
            start_column,
            end_column + 1,
        ):

            worksheet.cell(
                row=header_start_row,
                column=column,
            ).fill = period_fill

            worksheet.cell(
                row=header_start_row,
                column=column,
            ).border = header_border

    for index, column in enumerate(
        column_records
    ):

        excel_column = 5 + index

        cell = worksheet.cell(
            row=header_start_row + 1,
            column=excel_column,
            value=format_measure_label(
                column["measure"]
            ),
        )

        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        cell.fill = measure_fill
        cell.border = header_border

    worksheet.row_dimensions[
        header_start_row
    ].height = 22

    worksheet.row_dimensions[
        header_start_row + 1
    ].height = 22

    data_start_row = header_start_row + 2

    for row_index, (_, row) in enumerate(
        matrix_df.iterrows()
    ):

        excel_row = data_start_row + row_index

        metric_cell = worksheet.cell(
            row=excel_row,
            column=1,
            value=row["Metric"],
        )

        metric_cell.font = metric_font

        metric_cell.alignment = Alignment(
            vertical="center",
        )

        dimensions_cell = worksheet.cell(
            row=excel_row,
            column=2,
            value=row["Dimensions"],
        )

        dimensions_cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

        fx_basis_cell = worksheet.cell(
            row=excel_row,
            column=3,
            value=row["FX Basis"],
        )

        fx_basis_cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

        uom_cell = worksheet.cell(
            row=excel_row,
            column=4,
            value=row["UoM"],
        )

        uom_cell.alignment = Alignment(
            vertical="center",
        )

        for index, column in enumerate(
            column_records
        ):

            excel_column = 5 + index

            key = (
                column["period"],
                column["period_type"],
                column["measure"],
            )

            value = row.get(key)

            cell = worksheet.cell(
                row=excel_row,
                column=excel_column,
                value=value,
            )

            cell.alignment = Alignment(
                horizontal="right",
                vertical="center",
            )

            apply_value_number_format(
                cell,
                unit=row["UoM"],
                value=value,
            )

        for column in range(
            1,
            5 + len(column_records),
        ):

            worksheet.cell(
                row=excel_row,
                column=column,
            ).border = border

    worksheet.freeze_panes = f"E{data_start_row}"

    if not matrix_df.empty:

        worksheet.auto_filter.ref = (
            f"A{header_start_row}:"
            f"{get_column_letter(4 + len(column_records))}"
            f"{data_start_row + len(matrix_df) - 1}"
        )

    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 32
    worksheet.column_dimensions["C"].width = 18
    worksheet.column_dimensions["D"].width = 14

    for column in range(
        5,
        5 + len(column_records),
    ):

        worksheet.column_dimensions[
            get_column_letter(column)
        ].width = 15

    worksheet.sheet_view.showGridLines = False

    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0

    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    worksheet.print_title_rows = (
        f"{header_start_row}:"
        f"{header_start_row + 1}"
    )

    return worksheet


# ============================================================
# Create KPI Detail worksheet
# ============================================================

def create_kpi_detail_sheet(
    workbook,
    result_df,
):

    worksheet = workbook.create_sheet(
        "KPI Detail"
    )

    columns = [
        "Source Document",
        "Page",
        "Company",
        "Holding",
        "Reporting Period",
        "Metric",
        "Dimensions",
        "Period",
        "Period Type",
        "Measure",
        "Value",
        "UoM",
        "Currency",
        "FX Basis",
    ]

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    header_font = Font(
        bold=True,
    )

    thin = Side(
        style="thin",
        color="D9D9D9",
    )

    for column_index, column_name in enumerate(
        columns,
        start=1,
    ):

        cell = worksheet.cell(
            row=1,
            column=column_index,
            value=column_name,
        )

        cell.font = header_font
        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        cell.border = Border(
            bottom=thin,
        )

    if result_df.empty:
        return worksheet

    for row_index, (_, row) in enumerate(
        result_df.iterrows(),
        start=2,
    ):

        values = [
            row.get("source_document"),
            row.get("page_number"),
            row.get("company"),
            row.get("holding"),
            row.get("reporting_period"),
            row.get("kpi"),
            format_dimensions(
                row.get("dimensions")
            ),
            row.get("period"),
            row.get("period_type"),
            format_measure_label(
                row.get("measure")
            ),
            row.get("value"),
            row.get("unit"),
            row.get("currency"),
            row.get("fx_basis"),
        ]

        for column_index, value in enumerate(
            values,
            start=1,
        ):

            if is_missing(value):
                value = None

            cell = worksheet.cell(
                row=row_index,
                column=column_index,
                value=value,
            )

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

            if column_index == 11:

                apply_value_number_format(
                    cell,
                    unit=row.get("unit"),
                    value=value,
                )

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(columns))}"
        f"{len(result_df) + 1}"
    )

    widths = {
        "A": 32,
        "B": 10,
        "C": 25,
        "D": 25,
        "E": 20,
        "F": 30,
        "G": 35,
        "H": 15,
        "I": 15,
        "J": 15,
        "K": 18,
        "L": 14,
        "M": 14,
        "N": 18,
    }

    for column, width in widths.items():
        worksheet.column_dimensions[
            column
        ].width = width

    worksheet.sheet_view.showGridLines = False

    return worksheet


# ============================================================
# Create Source Notes worksheet
# ============================================================

def create_source_notes_sheet(
    workbook,
    result_df,
):

    worksheet = workbook.create_sheet(
        "Source Notes"
    )

    columns = [
        "Source Document",
        "Page",
        "Metric",
        "Dimensions",
        "Period",
        "Period Type",
        "Measure",
        "Value",
        "UoM",
        "Currency",
        "FX Basis",
    ]

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    header_font = Font(
        bold=True,
    )

    for column_index, column_name in enumerate(
        columns,
        start=1,
    ):

        cell = worksheet.cell(
            row=1,
            column=column_index,
            value=column_name,
        )

        cell.font = header_font
        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    if result_df.empty:
        return worksheet

    for row_index, (_, row) in enumerate(
        result_df.iterrows(),
        start=2,
    ):

        values = [
            row.get("source_document"),
            row.get("page_number"),
            row.get("kpi"),
            format_dimensions(
                row.get("dimensions")
            ),
            row.get("period"),
            row.get("period_type"),
            format_measure_label(
                row.get("measure")
            ),
            row.get("value"),
            row.get("unit"),
            row.get("currency"),
            row.get("fx_basis"),
        ]

        for column_index, value in enumerate(
            values,
            start=1,
        ):

            if is_missing(value):
                value = None

            cell = worksheet.cell(
                row=row_index,
                column=column_index,
                value=value,
            )

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

            if column_index == 8:

                apply_value_number_format(
                    cell,
                    unit=row.get("unit"),
                    value=value,
                )

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(columns))}"
        f"{len(result_df) + 1}"
    )

    widths = {
        "A": 32,
        "B": 10,
        "C": 30,
        "D": 35,
        "E": 15,
        "F": 15,
        "G": 15,
        "H": 18,
        "I": 14,
        "J": 14,
        "K": 18,
    }

    for column, width in widths.items():
        worksheet.column_dimensions[
            column
        ].width = width

    worksheet.sheet_view.showGridLines = False

    return worksheet


# ============================================================
# Build Excel workbook
# ============================================================

def create_excel_workbook(result_df):

    workbook = Workbook()

    default_sheet = workbook.active

    workbook.remove(
        default_sheet
    )

    create_financial_kpis_sheet(
        workbook,
        result_df,
    )

    create_kpi_detail_sheet(
        workbook,
        result_df,
    )

    create_source_notes_sheet(
        workbook,
        result_df,
    )

    workbook.active = (
        workbook.sheetnames.index(
            "Financial KPIs"
        )
    )

    output = BytesIO()

    workbook.save(output)

    archive_path = os.path.join(
        XLSX_DIRECTORY,
        "parsed_kpi.xlsx",
    )

    workbook.save(archive_path)

    output.seek(0)

    return output.getvalue()


# ============================================================
# Excel archive helpers
# ============================================================

def split_conflict_values(value):
    """
    Treat pipe-separated values as conflicting alternatives.
    """

    if is_missing(value):
        return []

    if isinstance(value, str):

        parts = [
            part.strip()
            for part in value.split("|")
        ]

        parts = [
            part
            for part in parts
            if part
        ]

        if len(parts) > 1:
            return list(
                dict.fromkeys(parts)
            )

    return [value]


def collapse_duplicate_pipe_values(value):
    """
    Collapse duplicate alternatives in a pipe-separated value.

    Examples:
        "28.8 | 28.8" -> "28.8"
        "28.8 | 29.1" -> "28.8 | 29.1"
    """

    if is_missing(value):
        return value

    if not isinstance(value, str) or "|" not in value:
        return value

    parts = [
        part.strip()
        for part in value.split("|")
        if part.strip()
    ]

    if len(parts) <= 1:
        return value

    unique_parts = list(
        dict.fromkeys(parts)
    )

    if len(unique_parts) == 1:
        return unique_parts[0]

    return " | ".join(unique_parts)


def display_cell_value(value):
    """Return a clean value suitable for Streamlit display."""

    if is_missing(value):
        return ""

    return collapse_duplicate_pipe_values(value)


def find_label_value(workbook, label):
    """Find a metadata label anywhere in the workbook."""

    target = str(label).strip().lower()

    for worksheet in workbook.worksheets:

        for row in worksheet.iter_rows():

            for cell in row:

                if (
                    isinstance(cell.value, str)
                    and cell.value.strip().lower()
                    == target
                ):

                    right_cell = worksheet.cell(
                        row=cell.row,
                        column=cell.column + 1,
                    )

                    return right_cell.value

    return None


def workbook_metadata(path):
    """Read Company and Reporting Period from workbook."""

    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    try:

        company = find_label_value(
            workbook,
            "Company",
        )

        reporting_period = find_label_value(
            workbook,
            "Reporting Period",
        )

        return {
            "company": company,
            "reporting_period": reporting_period,
            "filename": os.path.basename(path),
            "path": path,
        }

    finally:
        workbook.close()


def scan_xlsx_archive(directory):
    """Read every .xlsx file in the archive directory."""

    if not os.path.isdir(directory):
        return []

    paths = sorted(
        (
            os.path.join(directory, filename)
            for filename in os.listdir(directory)
            if filename.lower().endswith(".xlsx")
            and not filename.startswith("~$")
        ),
        key=lambda value: os.path.basename(value).lower(),
    )

    records = []

    for path in paths:

        try:

            records.append(
                workbook_metadata(path)
            )

        except Exception as exc:

            records.append(
                {
                    "company": None,
                    "reporting_period": None,
                    "filename": os.path.basename(path),
                    "path": path,
                    "error": str(exc),
                }
            )

    return records


def unique_dropdown_values(records, key):
    """Return unique values for a metadata dropdown."""

    values = set()

    for record in records:

        raw_value = record.get(key)

        for value in split_conflict_values(
            raw_value
        ):

            if not is_missing(value):
                values.add(
                    str(value).strip()
                )

    return sorted(
        values,
        key=lambda value: (
            period_sort_key(value)
            if key == "reporting_period"
            else value.lower()
        ),
    )


def locate_financial_kpis_header(worksheet):
    """Find the generated Financial KPIs static header row."""

    required = {
        "Metric",
        "Dimensions",
        "FX Basis",
        "UoM",
    }

    for row_index in range(
        1,
        min(worksheet.max_row, 40) + 1,
    ):

        values = {
            worksheet.cell(
                row=row_index,
                column=column_index,
            ).value
            for column_index in range(1, 5)
        }

        if required.issubset(
            {
                str(value).strip()
                for value in values
                if value is not None
            }
        ):

            return row_index

    return None


def get_financial_matrix_from_workbook(
    path,
    replacements=None,
):
    """
    Read the Financial KPIs worksheet and reconstruct
    the matrix, applying priority metric ordering and
    increasing dimensions length ordering.
    """

    replacements = replacements or {}

    workbook = load_workbook(
        filename=path,
        read_only=False,
        data_only=True,
    )

    try:

        if "Financial KPIs" not in workbook.sheetnames:
            return pd.DataFrame(), []

        worksheet = workbook[
            "Financial KPIs"
        ]

        header_row = locate_financial_kpis_header(
            worksheet
        )

        if header_row is None:
            return pd.DataFrame(), []

        data_start_row = header_row + 2

        static_headers = [
            worksheet.cell(
                row=header_row,
                column=column_index,
            ).value
            for column_index in range(1, 5)
        ]

        column_records = []

        current_period = None
        current_period_type = None

        reverse_measure = {
            "Actual": "actual",
            "Actual LY": "actual_ly",
            "Budget": "budget",
            "Forecast": "forecast",
        }

        for column_index in range(
            5,
            worksheet.max_column + 1,
        ):

            period_header = worksheet.cell(
                row=header_row,
                column=column_index,
            ).value

            measure_header = worksheet.cell(
                row=header_row + 1,
                column=column_index,
            ).value

            if period_header is not None:

                period_text = str(
                    period_header
                ).strip()

                if " – " in period_text:

                    (
                        current_period,
                        current_period_type,
                    ) = period_text.split(
                        " – ",
                        1,
                    )

                elif " - " in period_text:

                    (
                        current_period,
                        current_period_type,
                    ) = period_text.split(
                        " - ",
                        1,
                    )

                else:

                    current_period = period_text
                    current_period_type = ""

            if (
                current_period is None
                or measure_header is None
            ):
                continue

            measure_text = str(
                measure_header
            ).strip()

            measure = reverse_measure.get(
                measure_text,
                measure_text.lower().replace(
                    " ",
                    "_",
                ),
            )

            column_records.append(
                {
                    "column_index": column_index,
                    "period": current_period,
                    "period_type": current_period_type,
                    "measure": measure,
                    "measure_label": measure_text,
                }
            )

        rows = []
        row_sort_keys = []

        for row_number in range(
            data_start_row,
            worksheet.max_row + 1,
        ):

            row_values = []

            for column_index in range(
                1,
                worksheet.max_column + 1,
            ):

                key = (
                    "Financial KPIs",
                    row_number,
                    column_index,
                )

                if key in replacements:

                    value = replacements[key]

                else:

                    value = worksheet.cell(
                        row=row_number,
                        column=column_index,
                    ).value

                row_values.append(value)

            if not any(
                not is_missing(value)
                for value in row_values
            ):
                continue

            row = {
                static_headers[index]: row_values[index]
                for index in range(
                    min(
                        4,
                        len(static_headers),
                    )
                )
            }

            for record in column_records:

                value = row_values[
                    record["column_index"] - 1
                ]

                if (
                    isinstance(value, (int, float))
                    and not isinstance(value, bool)
                    and value == 0
                ):
                    value = None

                row[
                    (
                        record["period"],
                        record["period_type"],
                        record["measure"],
                    )
                ] = value

            rows.append(row)

            # Sort key logic: priority list -> name -> dimensions length -> fx_basis -> unit
            metric_name = row.get("Metric")
            metric_rank = metric_sort_rank(metric_name)

            dimensions_val = row.get("Dimensions")
            try:
                dims_list = json.loads(dimensions_val) if isinstance(dimensions_val, str) else normalize_dimensions(dimensions_val)
                dims_len = len(dims_list)
            except Exception:
                dims_len = 0

            unmatched_name_key = (
                str(metric_name).strip().lower()
                if metric_rank == len(METRIC_ORDER)
                else ""
            )

            row_sort_keys.append(
                (
                    metric_rank,
                    unmatched_name_key,
                    dims_len,
                    str(row.get("FX Basis") or "").lower(),
                    str(row.get("UoM") or "").lower(),
                )
            )

        if rows:
            ordered = sorted(
                zip(row_sort_keys, rows),
                key=lambda item: item[0],
            )
            rows = [item[1] for item in ordered]

        matrix_df = pd.DataFrame(rows)

        if not matrix_df.empty:

            ordered_columns = [
                column
                for column in [
                    "Metric",
                    "Dimensions",
                    "FX Basis",
                    "UoM",
                ]
                if column in matrix_df.columns
            ]

            ordered_columns.extend(
                [
                    (
                        record["period"],
                        record["period_type"],
                        record["measure"],
                    )
                    for record in column_records
                ]
            )

            matrix_df = matrix_df.reindex(
                columns=ordered_columns
            )

        return (
            matrix_df,
            column_records,
        )

    finally:
        workbook.close()


def find_all_conflicts(path):
    """Find pipe-separated conflicts in every worksheet."""

    workbook = load_workbook(
        filename=path,
        read_only=False,
        data_only=True,
    )

    conflicts = []

    try:

        for worksheet in workbook.worksheets:

            for row in worksheet.iter_rows():

                for cell in row:

                    alternatives = split_conflict_values(
                        cell.value
                    )

                    if len(alternatives) <= 1:
                        continue

                    conflicts.append(
                        {
                            "sheet": worksheet.title,
                            "row": cell.row,
                            "column": cell.column,
                            "coordinate": cell.coordinate,
                            "original": cell.value,
                            "options": alternatives,
                        }
                    )

    finally:
        workbook.close()

    return conflicts


# ============================================================
# Conflict resolution helpers
# ============================================================

def conflict_key(path, conflict):
    """
    Generate a stable ID for one workbook conflict.

    This allows a submitted conflict to remain resolved across
    Streamlit reruns.
    """

    raw = (
        f"{os.path.abspath(path)}|"
        f"{conflict['sheet']}|"
        f"{conflict['row']}|"
        f"{conflict['column']}"
    )

    return hashlib.sha256(
        raw.encode("utf-8")
    ).hexdigest()[:20]


def get_conflict_metric_name(path, conflict):
    """
    Return the Metric name for a Financial KPIs conflict.

    For Financial KPIs conflicts, the metric is taken directly
    from column A of the corresponding data row.
    """

    if conflict["sheet"] != "Financial KPIs":
        return conflict["sheet"]

    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    try:

        worksheet = workbook["Financial KPIs"]

        header_row = locate_financial_kpis_header(
            worksheet
        )

        if header_row is None:
            return "Financial KPIs"

        data_start_row = header_row + 2

        if conflict["row"] < data_start_row:
            return "Financial KPIs"

        metric = worksheet.cell(
            row=conflict["row"],
            column=1,
        ).value

        if is_missing(metric):
            return "Financial KPIs"

        return str(metric).strip()

    except Exception:
        return "Financial KPIs"

    finally:
        workbook.close()


def conflict_label(path, conflict):
    """
    Display only the metric name rather than sheet/cell/value
    information.
    """

    return get_conflict_metric_name(
        path,
        conflict,
    )


def render_conflict_resolution(path):
    """
    Show unresolved conflicts and let the user submit one
    selected value at a time.

    Once submitted, that conflict is persisted in session state,
    disappears from the alert, and the table is reconstructed
    using the submitted value.
    """

    conflicts = find_all_conflicts(path)

    if not conflicts:
        return {}

    if "conflict_resolutions" not in st.session_state:
        st.session_state["conflict_resolutions"] = {}

    all_resolutions = st.session_state[
        "conflict_resolutions"
    ]

    unresolved_conflicts = []

    for conflict in conflicts:

        key = conflict_key(
            path,
            conflict,
        )

        if key not in all_resolutions:
            unresolved_conflicts.append(
                conflict
            )

    if not unresolved_conflicts:
        return {
            (
                conflict["sheet"],
                conflict["row"],
                conflict["column"],
            ): all_resolutions[
                conflict_key(path, conflict)
            ]
            for conflict in conflicts
            if conflict_key(path, conflict)
            in all_resolutions
        }

    st.warning(
        f"This workbook contains "
        f"{len(unresolved_conflicts)} "
        "conflicting field value(s)."
    )

    replacements = {}

    with st.expander(
        "Resolve conflicting values",
        expanded=False,
    ):

        for conflict in unresolved_conflicts:

            conflict_id = conflict_key(
                path,
                conflict,
            )

            options = conflict["options"]

            metric_name = conflict_label(
                path,
                conflict,
            )

            select_col, button_col = st.columns(
                [3, 1],
                gap="medium",
            )

            with select_col:

                selected = st.selectbox(
                    metric_name,
                    options=options,
                    key=f"conflict_select_{conflict_id}",
                )

            with button_col:

                st.markdown(
                    "<div style='height: 28px'></div>",
                    unsafe_allow_html=True,
                )

                submitted = st.button(
                    "Submit",
                    key=f"conflict_submit_{conflict_id}",
                    type="primary",
                    width="content",
                )

            if submitted:

                all_resolutions[
                    conflict_id
                ] = selected

                st.session_state[
                    "conflict_resolutions"
                ] = all_resolutions

                st.rerun()

    for conflict in conflicts:

        conflict_id = conflict_key(
            path,
            conflict,
        )

        if conflict_id in all_resolutions:

            replacements[
                (
                    conflict["sheet"],
                    conflict["row"],
                    conflict["column"],
                )
            ] = all_resolutions[
                conflict_id
            ]

    return replacements


def matching_workbooks(
    records,
    company,
    reporting_period,
):
    """Find archive workbooks matching both selections."""

    matches = []

    for record in records:

        company_values = {
            str(value).strip()
            for value in split_conflict_values(
                record.get("company")
            )
            if not is_missing(value)
        }

        period_values = {
            str(value).strip()
            for value in split_conflict_values(
                record.get("reporting_period")
            )
            if not is_missing(value)
        }

        if (
            company in company_values
            and reporting_period in period_values
        ):
            matches.append(record)

    return matches


def render_financial_preview(
    matrix_df,
    column_records,
    selected_fx_basis=None,
):
    """Render the matrix in the same logical structure as Excel."""

    if matrix_df.empty:

        st.warning(
            "No Financial KPIs table was found in this workbook."
        )

        return

    preview_df = matrix_df.copy()

    # FX Basis is selected above the table on the homepage, so it
    # is no longer displayed as a table column.
    if (
        selected_fx_basis is not None
        and "FX Basis" in preview_df.columns
    ):
        preview_df = preview_df[
            preview_df["FX Basis"].astype(str).str.strip().str.lower()
            == str(selected_fx_basis).strip().lower()
        ].copy()

    if "FX Basis" in preview_df.columns:
        preview_df = preview_df.drop(
            columns=["FX Basis"]
        )

    renamed_columns = {}

    for column in preview_df.columns:

        if isinstance(
            column,
            tuple,
        ):

            period = column[0]
            period_type = column[1]
            measure = column[2]

            renamed_columns[column] = (
                f"{period} – "
                f"{period_type} – "
                f"{format_measure_label(measure)}"
            )

    preview_df = preview_df.rename(
        columns=renamed_columns
    )

    preview_df = preview_df.map(
        collapse_duplicate_pipe_values
    )

    preview_df = sanitize_for_streamlit(
        preview_df
    )

    st.dataframe(
        preview_df,
        width="stretch",
        hide_index=True,
    )


# ============================================================
# Homepage
# ============================================================

def render_homepage():

    archive_records = scan_xlsx_archive(
        XLSX_DIRECTORY
    )

    # Keep KPI browsing and parsing as separate top-level pages.
    kpi_tab, parsing_tab = st.tabs(
        [
            "KPI Tables",
            "Parsing",
        ]
    )

    with kpi_tab:

        if not archive_records:

            st.markdown(
                """
                <div class="vp-card">
                    <div class="vp-card-title">
                        Your workspace is ready.
                    </div>
                    <div class="vp-card-description">
                        No exported workbooks have been added yet.
                        Use the Parsing tab to upload financial
                        documents and create your first KPI workbook.
                    </div>
                </div>
                """,
                unsafe_allow_html=True,
            )

        else:

            broken_records = [
                record
                for record in archive_records
                if record.get("error")
            ]

            if broken_records:

                with st.expander(
                    f"⚠️ {len(broken_records)} workbook(s) could not be read"
                ):

                    for record in broken_records:

                        st.error(
                            f"**{record['filename']}** — "
                            f"{record['error']}"
                        )

            company_options = unique_dropdown_values(
                archive_records,
                "company",
            )

            period_options = unique_dropdown_values(
                archive_records,
                "reporting_period",
            )

            if not company_options:

                st.warning(
                    "No Company values were found in the exported "
                    "Excel files."
                )

            elif not period_options:

                st.warning(
                    "No Reporting Period values were found in the "
                    "exported Excel files."
                )

            else:

                st.markdown(
                    """
                    <div class="vp-section-header">
                        <div>
                            <div class="vp-section-title">
                                Performance archive
                            </div>
                            <div class="vp-section-copy">
                                Select a company and reporting period to
                                open a previously exported KPI workbook.
                            </div>
                        </div>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )

                col1, col2, col3 = st.columns(
                    3,
                    gap="large",
                )

                with col1:

                    selected_company = st.selectbox(
                        "Company",
                        options=company_options,
                        key="homepage_company",
                    )

                with col2:

                    selected_period = st.selectbox(
                        "Reporting Period",
                        options=period_options,
                        key="homepage_reporting_period",
                    )

                matches = matching_workbooks(
                    archive_records,
                    selected_company,
                    selected_period,
                )

                if not matches:

                    st.info(
                        "No exported workbook matches this Company and "
                        "Reporting Period selection."
                    )

                else:

                    if len(matches) > 1:

                        st.warning(
                            f"{len(matches)} workbooks match this selection."
                        )

                        workbook_labels = [
                            record["filename"]
                            for record in matches
                        ]

                        selected_filename = st.selectbox(
                            "Workbook",
                            options=workbook_labels,
                            key="homepage_workbook",
                        )

                        selected_record = next(
                            record
                            for record in matches
                            if record["filename"] == selected_filename
                        )

                    else:

                        selected_record = matches[0]

                    current_path = os.path.abspath(
                        selected_record["path"]
                    )

                    st.session_state[
                        "active_conflict_workbook"
                    ] = current_path

                    st.write("")

                    replacements = render_conflict_resolution(
                        selected_record["path"]
                    )

                    matrix_df, column_records = (
                        get_financial_matrix_from_workbook(
                            selected_record["path"],
                            replacements=replacements,
                        )
                    )

                    # FX Basis is a page-level selector, alongside
                    # Company and Reporting Period. The Financial KPIs
                    # table below is filtered to the selected basis and
                    # therefore does not display FX Basis as a column.
                    fx_basis_options = [
                        basis
                        for basis in FX_BASIS_OPTIONS
                        if (
                            not matrix_df.empty
                            and "FX Basis" in matrix_df.columns
                            and matrix_df["FX Basis"]
                            .fillna("")
                            .astype(str)
                            .str.strip()
                            .str.lower()
                            .eq(basis.lower())
                            .any()
                        )
                    ]

                    if not fx_basis_options:
                        fx_basis_options = FX_BASIS_OPTIONS.copy()

                    # Rendered into the third column of the selector
                    # row above, so it sits alongside Company and
                    # Reporting Period.
                    selected_fx_basis = col3.selectbox(
                        "FX Basis",
                        options=fx_basis_options,
                        key="homepage_fx_basis",
                    )

                    st.markdown(
                        """
                        <div class="vp-section-header">
                            <div>
                                <div class="vp-section-title">
                                    Financial KPIs
                                </div>
                            </div>
                        </div>
                        """,
                        unsafe_allow_html=True,
                    )

                    render_financial_preview(
                        matrix_df,
                        column_records,
                        selected_fx_basis=selected_fx_basis,
                    )

    with parsing_tab:
        render_extraction_ui()


# ============================================================
# Extraction UI
# ============================================================

def render_extraction_ui():

    st.markdown(
        """
        <div class="vp-card">
            <div class="vp-card-title">
                Extract financial KPIs
            </div>

            <div class="vp-card-description">
                Upload financial reports and VerePort will extract
                structured KPI observations including periods,
                measures, dimensions, currency and FX basis.
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    uploaded_files = st.file_uploader(
        "Choose financial documents",
        type=ALLOWED_EXTENSIONS,
        accept_multiple_files=True,
        key="extraction_file_uploader",
    )

    if not uploaded_files:
        return

    if len(uploaded_files) > MAX_FILES:

        st.error(
            f"You selected {len(uploaded_files)} files. "
            f"The maximum allowed is {MAX_FILES} files."
        )

        return

    oversized_files = [
        uploaded_file
        for uploaded_file in uploaded_files
        if uploaded_file.size > MAX_FILE_SIZE_BYTES
    ]

    if oversized_files:

        st.error(
            f"{len(oversized_files)} file(s) exceed the "
            f"{MAX_FILE_SIZE_MB} MB per-file limit."
        )

        for uploaded_file in oversized_files:

            size_mb = (
                uploaded_file.size
                / (1024 * 1024)
            )

            st.error(
                f"📄 **{uploaded_file.name}** — "
                f"{size_mb:.2f} MB "
                f"(maximum {MAX_FILE_SIZE_MB} MB)"
            )

        return

    for uploaded_file in uploaded_files:

        size_mb = (
            uploaded_file.size
            / (1024 * 1024)
        )

        st.markdown(
            f"""
            <div class="vp-meta-row">
                <div class="vp-file-chip">
                    📄 {uploaded_file.name}
                </div>

                <div class="vp-file-chip">
                    {size_mb:.2f} MB
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    st.write("")

    # Smaller Extract KPIs button.
    extract_button_col = st.columns(
        [1, 2, 1]
    )[0]

    with extract_button_col:

        extract_clicked = st.button(
            "Extract KPIs",
            type="primary",
            key="extract_kpis_button",
            width="content",
        )

    if not extract_clicked:
        render_extraction_results()
        return

    headers = {
        "Modal-Key": MODAL_KEY,
        "Modal-Secret": MODAL_SECRET,
    }

    total_files = len(uploaded_files)

    progress = st.progress(0)
    status_text = st.empty()
    results = []

    with st.spinner(
        f"Starting {total_files} extraction requests..."
    ):

        with ThreadPoolExecutor(
            max_workers=min(
                MAX_WORKERS,
                total_files,
            )
        ) as executor:

            futures = []

            for index, uploaded_file in enumerate(
                uploaded_files
            ):

                if index > 0:
                    time.sleep(
                        REQUEST_START_DELAY_SECONDS
                    )

                future = executor.submit(
                    extract_single_file,
                    uploaded_file,
                    headers,
                )

                futures.append(
                    (
                        future,
                        uploaded_file.name,
                    )
                )

                status_text.info(
                    f"Started {index + 1}/{total_files}: "
                    f"{uploaded_file.name}"
                )

            completed = 0

            for future, filename in futures:

                try:
                    result = future.result()

                except Exception as exc:

                    result = {
                        "filename": filename,
                        "success": False,
                        "response_data": None,
                        "error": (
                            "Unexpected worker error: "
                            f"{exc}"
                        ),
                    }

                results.append(result)

                completed += 1

                progress.progress(
                    completed / total_files
                )

                if result["success"]:

                    status_text.success(
                        f"Completed {completed}/{total_files}: "
                        f"{filename}"
                    )

                else:

                    status_text.error(
                        f"Failed {completed}/{total_files}: "
                        f"{filename}"
                    )

    all_rows = []
    all_skipped = []
    all_failed_extractions = []
    request_failures = []

    successful_files = 0
    total_extractions = 0

    for result in results:

        filename = result["filename"]

        if not result["success"]:

            request_failures.append(
                {
                    "file": filename,
                    "error": result["error"],
                }
            )

            continue

        successful_files += 1

        response_data = (
            result["response_data"]
            or {}
        )

        if not isinstance(
            response_data,
            dict,
        ):

            all_failed_extractions.append(
                {
                    "source": filename,
                    "error": (
                        "API response was not a JSON object."
                    ),
                }
            )

            continue

        extractions = (
            response_data.get("extractions")
            or []
        )

        if extractions:

            total_extractions += len(
                extractions
            )

        elif isinstance(
            response_data.get("kpis"),
            list,
        ):

            total_extractions += 1

        file_df = flatten_extractions(
            response_data,
            default_source=filename,
        )

        if not file_df.empty:
            all_rows.append(file_df)

        skipped = response_data.get(
            "skipped",
            [],
        )

        if skipped:

            if isinstance(
                skipped,
                list,
            ):

                all_skipped.extend(skipped)

            else:

                all_skipped.append(skipped)

        failed = response_data.get(
            "failed_extractions",
            [],
        )

        if failed:

            if isinstance(
                failed,
                list,
            ):

                all_failed_extractions.extend(
                    failed
                )

            else:

                all_failed_extractions.append(
                    failed
                )

    if all_rows:

        result_df = pd.concat(
            all_rows,
            ignore_index=True,
            sort=False,
        )

    else:

        result_df = pd.DataFrame()

    if not result_df.empty:

        st.session_state["result_df"] = result_df
        create_excel_workbook(result_df)

        st.markdown(
            f"""
            <div class="vp-result-banner">
                <div class="vp-result-title">
                    Extraction complete
                </div>

                <div class="vp-result-copy">
                    {successful_files}/{total_files}
                    files processed ·
                    {len(result_df):,} observations created.
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    else:

        st.session_state.pop(
            "result_df",
            None,
        )

        if successful_files:

            st.warning(
                f"Completed {successful_files}/{total_files} files, "
                "but no KPI observations were found."
            )

            if total_extractions == 0:

                st.info(
                    "The API returned no extraction items."
                )

            else:

                st.info(
                    f"The API returned {total_extractions} "
                    "extraction(s), but they contained no KPI "
                    "observations."
                )

        else:

            st.error(
                "All extraction requests failed."
            )

    if request_failures:

        with st.expander(
            f"Failed requests ({len(request_failures)})",
            expanded=True,
        ):

            for failure in request_failures:

                st.error(
                    f"**{failure['file']}** — "
                    f"{failure['error']}"
                )

    if all_skipped:

        with st.expander(
            f"Skipped items ({len(all_skipped)})"
        ):

            st.json(all_skipped)

    ## don't show failed extractions in web view
    # if all_failed_extractions:

    #     with st.expander(
    #         f"Failed extractions "
    #         f"({len(all_failed_extractions)})"
    #     ):

    #         st.json(all_failed_extractions)

    render_extraction_results()


# ============================================================
# Extraction results
# ============================================================

def render_extraction_results():

    if "result_df" not in st.session_state:
        return

    result_df = st.session_state["result_df"]

    st.divider()

    st.markdown(
        """
        <div class="vp-section-header">
            <div>
                <div class="vp-section-title">
                    Extraction overview
                </div>

                <div class="vp-section-copy">
                    A concise view of the structured financial
                    observations generated from your documents.
                </div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    col1, col2, col3 = st.columns(
        3,
        gap="medium",
    )

    with col1:

        st.metric(
            "Sources",
            result_df["source"].nunique(),
        )

    with col2:

        st.metric(
            "KPIs",
            result_df["kpi"].nunique(),
        )

    with col3:

        st.metric(
            "Observations",
            len(result_df),
        )

    st.divider()

    st.markdown(
        """
        <div class="vp-section-header">
            <div>
                <div class="vp-section-title">
                    Financial KPI preview
                </div>

                <div class="vp-section-copy">
                    KPIs are presented as rows, with dimensions,
                    FX Basis and UoM as row identifiers.
                </div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    matrix_df, column_records = (
        build_financial_matrix(
            result_df
        )
    )

    if not matrix_df.empty:

        preview_df = matrix_df.copy()

        renamed_columns = {}

        for column in preview_df.columns:

            if isinstance(
                column,
                tuple,
            ):

                period = column[0]
                period_type = column[1]
                measure = column[2]

                renamed_columns[column] = (
                    f"{period} – "
                    f"{period_type} – "
                    f"{format_measure_label(measure)}"
                )

        preview_df = preview_df.rename(
            columns=renamed_columns
        )

        preview_df = preview_df.map(
            collapse_duplicate_pipe_values
        )

        preview_df = sanitize_for_streamlit(
            preview_df
        )

        st.dataframe(
            preview_df,
            width="stretch",
            hide_index=True,
        )

    else:

        st.warning(
            "No valid financial observations found."
        )

    st.divider()

    st.markdown(
        """
        <div class="vp-section-header">
            <div>
                <div class="vp-section-title">
                    Export
                </div>

                <div class="vp-section-copy">
                    Download a formatted Excel workbook containing
                    the financial matrix, underlying KPI detail and
                    source traceability.
                </div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    try:

        excel_data = create_excel_workbook(
            result_df
        )

        st.download_button(
            label="Download Financial KPI Excel",
            data=excel_data,
            file_name="financial_kpis.xlsx",
            mime=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
            type="primary",
            width="stretch",
        )

    except Exception as exc:

        st.error(
            f"Could not create Excel export: {exc}"
        )

    st.divider()

    with st.expander(
        "KPI Detail",
        expanded=False,
    ):

        detail_preview = result_df.copy()

        if "dimensions" in detail_preview.columns:

            detail_preview["dimensions"] = (
                detail_preview[
                    "dimensions"
                ].apply(
                    format_dimensions
                )
            )

        display_columns = [
            "source_document",
            "page_number",
            "kpi",
            "dimensions",
            "fx_basis",
            "period",
            "period_type",
            "measure",
            "value",
            "unit",
            "currency",
        ]

        available_columns = [
            column
            for column in display_columns
            if column in detail_preview.columns
        ]

        detail_preview = detail_preview[
            available_columns
        ].copy()

        detail_preview = detail_preview.rename(
            columns={
                "source_document": "Source Document",
                "page_number": "Page",
                "kpi": "Metric",
                "dimensions": "Dimensions",
                "fx_basis": "FX Basis",
                "period": "Period",
                "period_type": "Period Type",
                "measure": "Measure",
                "value": "Value",
                "unit": "UoM",
                "currency": "Currency",
            }
        )

        detail_preview = sanitize_for_streamlit(
            detail_preview
        )

        st.dataframe(
            detail_preview,
            width="stretch",
            hide_index=True,
        )


# ============================================================
# App footer
# ============================================================

def render_footer():

    st.markdown(
        """
        <div class="vp-footer">
            VerePort · Financial intelligence workspace
        </div>
        """,
        unsafe_allow_html=True,
    )


# ============================================================
# Run application
# ============================================================

render_homepage()

render_footer()